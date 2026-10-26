from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.services.excel_report import build_workbook, flujo_por_hora

DATA = {
    "flujo": 3, "permanencia": 25.0, "ocupacion": 1, "tracks": 2, "detecciones": 40,
    "zonas": [{"id": 1, "zona": "Entrada", "flujo": 2, "permanencia": 20.0, "ocupacion": 0},
              {"id": 2, "zona": "Cajas", "flujo": 1, "permanencia": 35.0, "ocupacion": 1}],
    "advanced": {"confidence_avg": 81.5, "completed_visits": 2, "dwell_max": 35.0, "zones_count": 2,
                 "top_flow": "Entrada", "top_flow_value": 2, "top_dwell": "Cajas", "top_dwell_value": 35.0},
}
EVENTOS = [
    {"id_evento": 1, "track": 7, "zona": "Entrada", "entrada": datetime(2026, 10, 1, 10, 5),
     "salida": datetime(2026, 10, 1, 10, 5, 20), "permanencia_seg": 20.0},
    {"id_evento": 2, "track": 8, "zona": "Entrada", "entrada": datetime(2026, 10, 1, 11, 0),
     "salida": None, "permanencia_seg": 0},
    {"id_evento": 3, "track": 7, "zona": "Cajas", "entrada": datetime(2026, 10, 1, 11, 10),
     "salida": datetime(2026, 10, 1, 11, 10, 35), "permanencia_seg": 35.0},
]
META = {"tienda": "Tienda Demo", "generado_por": "Analista", "generado": datetime(2026, 10, 13, 9, 0),
        "id_reporte": 5, "store_id": 1, "modelo": "yolo11n.pt", "tracker": "bytetrack.yaml",
        "conf": 0.35, "iou": 0.7, "eventos_exportados": 3}


def _libro(data=DATA, eventos=EVENTOS):
    return load_workbook(BytesIO(build_workbook(data, eventos, META)))


def test_hojas_del_reporte():
    assert _libro().sheetnames == ["Resumen", "Zonas", "Flujo por hora", "Eventos", "Trazabilidad"]


def test_zonas_y_totales_como_formulas():
    ws = _libro()["Zonas"]
    assert [ws["A2"].value, ws["B2"].value, ws["A3"].value] == ["Entrada", 2, "Cajas"]
    assert ws["B4"].value == "=SUM(B2:B3)"
    assert ws["E2"].value == "=IF($B$4=0,0,B2/$B$4)"
    assert ws.freeze_panes == "B2" and ws.auto_filter.ref == "A1:E3"
    assert len(ws._charts) == 1


def test_kpis_del_resumen_referencian_otras_hojas():
    ws = _libro()["Resumen"]
    formulas = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 30)}
    assert formulas["Flujo total"] == "=Zonas!B4"
    assert formulas["Permanencia media"] == '=IFERROR(AVERAGEIF(Eventos!F2:F4,">0"),0)'
    assert formulas["Visitas completadas"] == "=COUNT(Eventos!E2:E4)"
    assert formulas["Confianza YOLO promedio"] == 0.815


def test_flujo_por_hora_separa_fechas_y_zonas():
    assert flujo_por_hora(EVENTOS) == {"2026-10-01 10:00": {"Entrada": 1},
                                       "2026-10-01 11:00": {"Entrada": 1, "Cajas": 1}}


def test_reporte_sin_datos_no_falla():
    vacio = {**DATA, "flujo": 0, "zonas": [], "advanced": {**DATA["advanced"], "top_flow": "Sin datos"}}
    ws = _libro(vacio, [])["Resumen"]
    valores = [ws.cell(row=r, column=2).value for r in range(1, 30)]
    assert "Sin datos" in valores

def test_todas_las_hojas_tienen_pie_de_pagina():
    libro = _libro()

    for hoja in libro.worksheets:
        assert hoja.oddFooter.center.text == "SmartRetail Vision"

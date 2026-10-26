"""
UC007 · Exportación del reporte a Excel (.xlsx): equivalente en Python de Apache POI.

Equivalencias (Apache POI / XSSF en Java  ->  openpyxl en Python):
    new XSSFWorkbook()                      -> Workbook()
    wb.createSheet("Zonas")                 -> wb.create_sheet("Zonas")
    row.createCell(c).setCellValue(v)       -> ws.cell(row=r, column=c, value=v)
    cell.setCellFormula("SUM(B2:B9)")       -> ws["B10"] = "=SUM(B2:B9)"
    CellStyle, XSSFFont, FillPattern        -> Font, PatternFill, Border, Alignment
    DataFormat "0.0" / "yyyy-mm-dd hh:mm"   -> cell.number_format
    sheet.createFreezePane(0, 1)            -> ws.freeze_panes = "A2"
    sheet.setAutoFilter(...)                -> ws.auto_filter.ref = "A1:F20"
    SheetConditionalFormatting              -> ws.conditional_formatting.add(..., ColorScaleRule)
    XSSFChart (barras / líneas)             -> openpyxl.chart.BarChart / LineChart
    wb.write(OutputStream)                  -> wb.save(BytesIO())

Reutiliza `report_data()` (la misma fuente que la vista /reports y el CSV),
por lo que los tres formatos muestran las mismas cifras. Los totales y KPIs se
escriben como fórmulas de Excel para que el analista pueda auditarlos.
"""
from __future__ import annotations

import json
import logging
from collections import defaultdict
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger("smartretail.reportes")

MAX_EVENTOS = 50_000  # Excel admite ~1 millón de filas; se limita para mantener el archivo liviano
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Estilos reutilizables (equivalente a crear un CellStyle por tipo de celda en POI).
# Colores tomados del tema del prototipo (app/web/static/css/app.css).
VERDE = "08734A"
VERDE_OSCURO = "092C20"
VERDE_CLARO = "DFF2E9"
GRIS = "F4F7F5"
FUENTE = "Arial"
_linea = Side(style="thin", color="C9D6CE")
BORDE = Border(left=_linea, right=_linea, top=_linea, bottom=_linea)
F_NORMAL = Font(name=FUENTE, size=10)
F_NEGRITA = Font(name=FUENTE, size=10, bold=True)
F_CABECERA = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
F_TITULO = Font(name=FUENTE, size=16, bold=True, color=VERDE_OSCURO)
F_NOTA = Font(name=FUENTE, size=9, italic=True, color="5F6F66")
R_CABECERA = PatternFill("solid", fgColor=VERDE)
R_TOTAL = PatternFill("solid", fgColor=VERDE_CLARO)
R_ETIQUETA = PatternFill("solid", fgColor=GRIS)

FMT_ENTERO = "#,##0"
FMT_DECIMAL = "#,##0.0"
FMT_FECHA = "yyyy-mm-dd hh:mm:ss"


# --------------------------------------------------------------------------- utilidades
def _cabecera(ws, fila: int, titulos: list[str]) -> None:
    for col, texto in enumerate(titulos, start=1):
        c = ws.cell(row=fila, column=col, value=texto)
        c.font, c.fill, c.border = F_CABECERA, R_CABECERA, BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila].height = 30


def _celda(ws, fila: int, col: int, valor, fmt: str | None = None, negrita: bool = False):
    c = ws.cell(row=fila, column=col, value=valor)
    c.font = F_NEGRITA if negrita else F_NORMAL
    c.border = BORDE
    if fmt:
        c.number_format = fmt
    return c


def _anchos(ws, anchos: list[float]) -> None:
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _nota(ws, fila: int, texto: str, ultima_col: int = 4) -> None:
    ws.cell(row=fila, column=1, value=texto).font = F_NOTA
    ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima_col)


def flujo_por_hora(eventos: list[dict]) -> dict[str, dict[str, int]]:
    """Cuenta ingresos (ENTRY) por franja horaria y zona. La franja incluye la
    fecha para no mezclar la misma hora de días distintos."""
    tabla: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    for e in eventos:
        tabla[e["entrada"].strftime("%Y-%m-%d %H:00")][e["zona"]] += 1
    return {h: dict(z) for h, z in tabla.items()}


# --------------------------------------------------------------------------- hojas
def _hoja_zonas(wb: Workbook, zonas: list[dict]) -> tuple[int, int]:
    ws = wb.create_sheet("Zonas")
    _cabecera(ws, 1, ["Zona", "Flujo (eventos de ingreso)", "Permanencia promedio (s)",
                      "Ocupación actual (presencias abiertas)", "% del flujo total"])
    ini, fin = 2, 1 + len(zonas)
    tot = fin + 1
    for i, z in enumerate(zonas):
        f = ini + i
        _celda(ws, f, 1, z["zona"])
        _celda(ws, f, 2, z["flujo"], FMT_ENTERO)
        _celda(ws, f, 3, z["permanencia"], FMT_DECIMAL)
        _celda(ws, f, 4, z["ocupacion"], FMT_ENTERO)
        _celda(ws, f, 5, f"=IF($B${tot}=0,0,B{f}/$B${tot})", "0.0%")
    if zonas:
        _celda(ws, tot, 1, "TOTAL", negrita=True)
        _celda(ws, tot, 2, f"=SUM(B{ini}:B{fin})", FMT_ENTERO, True)
        _celda(ws, tot, 3, "", negrita=True)
        _celda(ws, tot, 4, f"=SUM(D{ini}:D{fin})", FMT_ENTERO, True)
        _celda(ws, tot, 5, f"=SUM(E{ini}:E{fin})", "0.0%", True)
        for col in range(1, 6):
            ws.cell(row=tot, column=col).fill = R_TOTAL
        ws.auto_filter.ref = f"A1:E{fin}"
        ws.conditional_formatting.add(
            f"C{ini}:C{fin}",
            ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F8CBAD"))
        g = BarChart()
        g.type, g.title, g.legend = "bar", "Flujo por zona (eventos de ingreso)", None
        g.add_data(Reference(ws, min_col=2, min_row=1, max_row=fin), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=ini, max_row=fin))
        g.height, g.width = 7.5, 16
        ws.add_chart(g, f"A{tot + 3}")
    else:
        _nota(ws, 2, "Sin datos procesados: ejecute YOLO Live con zonas activas.", 5)
    ws.freeze_panes = "B2"
    _anchos(ws, [26, 16, 16, 20, 12])
    return ini, fin


def _hoja_horas(wb: Workbook, horas: dict[str, dict[str, int]], nombres_zona: list[str]) -> tuple[int, int, int]:
    ws = wb.create_sheet("Flujo por hora")
    col_total = len(nombres_zona) + 2
    _cabecera(ws, 1, ["Franja horaria (UTC)"] + nombres_zona + ["Total"])
    franjas = sorted(horas)
    for i, franja in enumerate(franjas):
        f = 2 + i
        _celda(ws, f, 1, franja)
        for j, zona in enumerate(nombres_zona):
            _celda(ws, f, 2 + j, horas[franja].get(zona, 0), FMT_ENTERO)
        _celda(ws, f, col_total, f"=SUM(B{f}:{get_column_letter(col_total - 1)}{f})", FMT_ENTERO, True)
    fin = 1 + len(franjas)
    if franjas:
        g = LineChart()
        g.title, g.legend = "Ingresos por franja horaria", None
        g.y_axis.title, g.x_axis.title = "Eventos", "Franja"
        g.add_data(Reference(ws, min_col=col_total, min_row=1, max_row=fin), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=1, min_row=2, max_row=fin))
        g.series[0].smooth = False  # líneas rectas: no se interpolan valores inexistentes
        g.height, g.width = 7.5, 18
        ws.add_chart(g, f"A{fin + 3}")
    else:
        _nota(ws, 2, "Sin eventos de ingreso registrados.", max(2, col_total))
    ws.freeze_panes = "B2"
    _anchos(ws, [20] + [14] * len(nombres_zona) + [10])
    return 2, fin, col_total


def _hoja_eventos(wb: Workbook, eventos: list[dict], truncado: bool) -> int:
    ws = wb.create_sheet("Eventos")
    _cabecera(ws, 1, ["ID evento", "ID track (temporal, anónimo)", "Zona",
                      "Entrada (UTC)", "Salida (UTC)", "Permanencia (s)"])
    for i, e in enumerate(eventos):
        f = 2 + i
        _celda(ws, f, 1, e["id_evento"])
        _celda(ws, f, 2, e["track"])
        _celda(ws, f, 3, e["zona"])
        _celda(ws, f, 4, e["entrada"], FMT_FECHA)
        _celda(ws, f, 5, e["salida"], FMT_FECHA)
        _celda(ws, f, 6, round(e["permanencia_seg"] or 0, 1), FMT_DECIMAL)
    fin = 1 + len(eventos)
    if eventos:
        ws.auto_filter.ref = f"A1:F{fin}"
    if truncado:
        _nota(ws, fin + 2, f"Se exportaron los {MAX_EVENTOS:,} eventos más recientes.", 6)
    ws.freeze_panes = "A2"
    _anchos(ws, [11, 16, 24, 20, 20, 15])
    return fin


def _hoja_resumen(ws, data: dict, meta: dict, zr: tuple[int, int], hr: tuple[int, int, int],
                  fin_eventos: int, truncado: bool) -> None:
    zi, zf = zr
    hi, hf, ht = hr
    a = data["advanced"]
    hay_zonas, hay_horas, hay_eventos = zf >= zi, hf >= hi, fin_eventos >= 2
    lt = get_column_letter(ht)

    ws["A1"] = "SmartRetail Vision · Reporte ejecutivo de analítica"
    ws["A1"].font = F_TITULO
    ws.merge_cells("A1:D1")
    if meta.get("datos_de_prueba"):
        ws["A2"] = "DATOS DE PRUEBA: generados para validar el reporte; no son resultados experimentales."
        ws["A2"].font = Font(name=FUENTE, size=10, bold=True, color="C00000")
        ws.merge_cells("A2:D2")
    fila = 3
    for etiqueta, valor in [("Tienda", meta["tienda"]), ("Generado por", meta["generado_por"]),
                            ("Fecha de generación", meta["generado"].strftime("%Y-%m-%d %H:%M:%S")),
                            ("ID de reporte", meta.get("id_reporte", "-"))]:
        _celda(ws, fila, 1, etiqueta, negrita=True).fill = R_ETIQUETA
        _celda(ws, fila, 2, valor)
        ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
        fila += 1

    fila += 1
    _cabecera(ws, fila, ["Indicador", "Valor", "Unidad", "Cómo se obtiene"])
    usar_eventos = hay_eventos and not truncado
    ev_f = f"Eventos!F2:F{fin_eventos}"
    kpis = [
        ("Flujo total", f"=Zonas!B{zf + 1}" if hay_zonas else 0, FMT_ENTERO, "eventos de zona",
         "Suma de la hoja Zonas"),
        ("Permanencia media", f'=IFERROR(AVERAGEIF({ev_f},">0"),0)' if usar_eventos else data["permanencia"],
         FMT_DECIMAL, "segundos", "Promedio de permanencias > 0 (hoja Eventos)"),
        ("Permanencia máxima", f"=MAX({ev_f})" if usar_eventos else a["dwell_max"], FMT_DECIMAL, "segundos",
         "Mayor permanencia registrada"),
        ("Ocupación actual", f"=Zonas!D{zf + 1}" if hay_zonas else 0, FMT_ENTERO, "personas",
         "Presencias abiertas (sin salida)"),
        ("Visitas completadas", f"=COUNT(Eventos!E2:E{fin_eventos})" if usar_eventos else a["completed_visits"],
         FMT_ENTERO, "visitas", "Eventos con salida registrada"),
        ("Tracks", data["tracks"], FMT_ENTERO, "IDs temporales", "Registros de la tabla track"),
        ("Detecciones YOLO", data["detecciones"], FMT_ENTERO, "observaciones", "Registros de la tabla deteccion"),
        ("Confianza YOLO promedio", a["confidence_avg"] / 100, "0.0%", "%",
         "Promedio de confidence; no equivale a mAP"),
        ("Zona con mayor flujo",
         f'=IFERROR(INDEX(Zonas!A{zi}:A{zf},MATCH(MAX(Zonas!B{zi}:B{zf}),Zonas!B{zi}:B{zf},0)),"Sin datos")'
         if hay_zonas else "Sin datos", None, "zona", "Máximo de la columna Flujo"),
        ("Zona con mayor permanencia",
         f'=IFERROR(INDEX(Zonas!A{zi}:A{zf},MATCH(MAX(Zonas!C{zi}:C{zf}),Zonas!C{zi}:C{zf},0)),"Sin datos")'
         if hay_zonas else "Sin datos", None, "zona", "Máximo de la columna Permanencia"),
        ("Hora pico",
         f"=IFERROR(INDEX('Flujo por hora'!A{hi}:A{hf},MATCH(MAX('Flujo por hora'!{lt}{hi}:{lt}{hf}),"
         f"'Flujo por hora'!{lt}{hi}:{lt}{hf},0)),\"Sin datos\")" if hay_horas else "Sin datos",
         None, "franja UTC", "Franja con más ingresos"),
    ]
    for nombre, valor, fmt, unidad, como in kpis:
        fila += 1
        _celda(ws, fila, 1, nombre, negrita=True)
        _celda(ws, fila, 2, valor, fmt, True).alignment = Alignment(horizontal="right")
        _celda(ws, fila, 3, unidad)
        _celda(ws, fila, 4, como)
    fila += 2
    _nota(ws, fila, "Privacidad: sin reconocimiento facial; los tracks son identificadores temporales anónimos.")
    if truncado:
        _nota(ws, fila + 1, "Eventos truncados: permanencias y visitas se toman del cálculo del sistema.")
    _anchos(ws, [28, 22, 16, 44])


def _hoja_trazabilidad(wb: Workbook, meta: dict) -> None:
    ws = wb.create_sheet("Trazabilidad")
    _cabecera(ws, 1, ["Parámetro", "Valor"])
    filas = [("ID de reporte", meta.get("id_reporte", "-")),
             ("Tienda (id)", meta.get("store_id")),
             ("Modelo detector", meta.get("modelo")),
             ("Tracker", meta.get("tracker")),
             ("Umbral de confianza YOLO", meta.get("conf")),
             ("Umbral IoU", meta.get("iou")),
             ("Eventos exportados", meta.get("eventos_exportados")),
             ("Fuente de datos", "report_data() · tablas zona, evento_zona, track, deteccion"),
             ("Zona horaria", "UTC (las marcas se guardan con datetime.utcnow)")]
    for i, (k, v) in enumerate(filas, start=2):
        _celda(ws, i, 1, k, negrita=True).fill = R_ETIQUETA
        _celda(ws, i, 2, v)
    _anchos(ws, [28, 60])


# --------------------------------------------------------------------------- API pública
def build_workbook(data: dict, eventos: list[dict], meta: dict, truncado: bool = False) -> bytes:
    """Arma el libro a partir de los datos de report_data() y devuelve sus bytes."""
    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"
    nombres = [z["zona"] for z in data["zonas"]]
    zr = _hoja_zonas(wb, data["zonas"])
    hr = _hoja_horas(wb, flujo_por_hora(eventos), nombres)
    fin_ev = _hoja_eventos(wb, eventos, truncado)
    _hoja_trazabilidad(wb, meta)
    _hoja_resumen(resumen, data, meta, zr, hr, fin_ev, truncado)
    for ws in wb.worksheets:  # impresión: horizontal y ajustada al ancho (PrintSetup en POI)
        ws.oddFooter.center.text = "SmartRetail Vision"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.properties.creator = "SmartRetail Vision"
    wb.properties.title = f"Reporte {meta['tienda']}"
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def make_xlsx(db, user, store_id=None, processing_id=None):
    """Equivalente de make_csv() para Excel: consulta, arma el libro y registra el Reporte."""
    from app.core.config import settings
    from app.models.entities import EventoZona, Reporte, Tienda, Track, Zona
    from app.services.reports import report_data

    data = report_data(db, store_id, processing_id)
    zone_ids = [z["id"] for z in data["zonas"]]
    eventos: list[dict] = []
    if zone_ids:
        query = (db.query(EventoZona, Zona.nombre, Track.external_track_id)
                 .join(Zona, Zona.id_zona == EventoZona.id_zona)
                 .join(Track, Track.id_track == EventoZona.id_track)
                 .filter(EventoZona.id_zona.in_(zone_ids))
                 .order_by(EventoZona.entrada.desc()))
        if processing_id is not None:
            query = query.filter(Track.id_procesamiento == processing_id)
        filas = query.limit(MAX_EVENTOS + 1).all()
        eventos = [{"id_evento": ev.id_evento, "track": tid, "zona": nombre, "entrada": ev.entrada,
                    "salida": ev.salida, "permanencia_seg": ev.permanencia_seg} for ev, nombre, tid in filas]
    truncado = len(eventos) > MAX_EVENTOS
    eventos = sorted(eventos[:MAX_EVENTOS], key=lambda e: e["entrada"])

    tienda = db.get(Tienda, store_id) if store_id else None
    rep = Reporte(id_usuario=user.id_usuario, tipo="XLSX_ANALYTICS",
                  parametros=json.dumps({"store_id": store_id, "processing_id": processing_id, "eventos": len(eventos), "truncado": truncado}))
    db.add(rep)
    db.flush()  # obtiene id_reporte para imprimirlo en el libro
    meta = {"tienda": tienda.nombre if tienda else "Todas las tiendas", "store_id": store_id, "processing_id": processing_id,
            "generado_por": user.nombre, "generado": datetime.now(), "id_reporte": rep.id_reporte,
            "modelo": settings.yolo_model, "tracker": settings.yolo_tracker,
            "conf": settings.yolo_confidence, "iou": settings.yolo_iou, "eventos_exportados": len(eventos)}
    contenido = build_workbook(data, eventos, meta, truncado)
    db.commit()
    log.info("Reporte XLSX #%s generado: %d zonas, %d eventos%s, %.1f KB", rep.id_reporte,
             len(zone_ids), len(eventos), " (truncado)" if truncado else "", len(contenido) / 1024)
    return contenido, rep
